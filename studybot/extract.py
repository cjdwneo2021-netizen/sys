import json
from pathlib import Path
from zipfile import ZipFile

from .storage import contained

PROMPT_VERSION = "lecture-v1"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def spreadsheet_rows(path):
    """Read cells and saved formula results without executing or updating a workbook."""
    from openpyxl import load_workbook
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError("Expanded spreadsheet exceeds 50 MiB; split the workbook")
    # File objects also handle Drive copies whose names no longer end in .xlsx.
    with path.open("rb") as formulas_file, path.open("rb") as values_file:
        formulas = load_workbook(formulas_file, read_only=True, data_only=False, keep_links=False)
        values = None
        try:
            values = load_workbook(values_file, read_only=True, data_only=True, keep_links=False)
            for sheet_index, sheet in enumerate(formulas.worksheets, 1):
                if (sheet.max_row or 1) * (sheet.max_column or 1) > 200_000:
                    raise ValueError("Spreadsheet grid exceeds 200,000 cells; split the worksheet")
                saved_rows = values[sheet.title].iter_rows()
                for row in sheet.iter_rows():
                    saved_row = next(saved_rows, ())
                    cells = []
                    for position, cell in enumerate(row):
                        if cell.value is None:
                            continue
                        text = str(cell.value)
                        if cell.data_type == "f":
                            cached = saved_row[position].value if position < len(saved_row) else None
                            text += (f" [저장된 결과: {cached}]" if cached is not None else
                                     " [저장된 결과 없음; 수식을 재계산하지 않음]")
                        elif cell.data_type == "e":
                            text += " [Excel 오류]"
                        cells.append(f"{cell.coordinate}: {text}")
                    if cells:
                        row_number = next(cell.row for cell in row if cell.value is not None)
                        yield sheet_index, sheet.title, row_number, "\t".join(cells)
        finally:
            formulas.close()
            if values is not None:
                values.close()

def extract_sources(root, manifest):
    """Never execute source code. Cell order and textual error output are preserved."""
    result = []
    for entry in manifest["files"]:
        if entry.get("status") != "downloaded":
            continue
        path = contained(root, entry["path"])
        suffix = path.suffix.lower()
        def add(source_id, label, text):
            if text.strip():
                result.append({"id": source_id, "label": label, "path": entry["path"], "text": text})
        if suffix == ".ipynb":
            notebook = json.loads(path.read_text("utf-8-sig"))
            for index, cell in enumerate(notebook.get("cells", []), 1):
                if cell.get("cell_type") not in ("code", "markdown", "raw"):
                    continue
                source_id = f'{entry["id"]}:cell-{index}'
                source = cell.get("source", "")
                source = "".join(source) if isinstance(source, list) else source
                add(source_id, f'{entry["name"]} / cell {index} ({cell["cell_type"]})', source)
                for position, output in enumerate(cell.get("outputs", []), 1):
                    text = output.get("text") or output.get("data", {}).get("text/plain")
                    if output.get("output_type") == "error":
                        text = f'{output.get("ename", "")}: {output.get("evalue", "")}'
                    if isinstance(text, list):
                        text = "".join(text)
                    if text:
                        if len(text) > 1500:
                            text = text[:1500] + "\n[나머지 실행 출력은 원본에 보관]"
                        add(f"{source_id}:output-{position}",
                            f'{entry["name"]} / cell {index} output {position}', text)
        elif suffix in (".txt", ".md", ".py"):
            add(entry["id"], entry["name"], path.read_text("utf-8-sig"))
        elif suffix == ".xlsx" or entry.get("mime_type") == XLSX_MIME:
            for sheet_index, title, row, text in spreadsheet_rows(path):
                add(f'{entry["id"]}:sheet-{sheet_index}:row-{row}',
                    f'{entry["name"]} / {title} / row {row}', text)
    return result

def chunks(sources, limit=6000):
    result, current = [], ""
    for source in sources:
        prefix = f'\n[SOURCE {source["id"]}] {source["label"]}\n'
        capacity = limit - len(prefix) - 1
        if capacity < 100:
            raise ValueError("Source label too long for chunk")
        for start in range(0, len(source["text"]), capacity):
            part = prefix + source["text"][start:start + capacity]
            if current and len(current) + len(part) > limit:
                result.append(current)
                current = ""
            current += part
    if current:
        result.append(current)
    return result
